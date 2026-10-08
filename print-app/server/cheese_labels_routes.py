"""
Routes for the cheese-labels batch printing flow.

This is a self-contained Flask Blueprint — it does its own BarTender COM
calls directly (win32com.client + pythoncom), the same way webhook_receiver.py
does for the standalone print.html tool. It does NOT go through /webhook,
and does not require HMAC signing — this is an internal-only tool used
directly from the factory floor, not called by JP Quality.

SETUP — the one thing you must edit before this works:
  CHEESE_LABEL_TEMPLATE_PATH below must point to the real .btw file for the
  cheese carton label — its Named Data Sources must be exactly (no spaces,
  this BarTender version doesn't allow them in data source names):
  מספר_אצווה, מזהה_קרטון, ברקוד, משקל, שם_אנגלית, שם_עברית).

Everything else (Excel reading, blueprint registration) needs no editing.
"""

import io
import logging
import pythoncom
import win32com.client
from flask import Blueprint, request, jsonify, Response
import openpyxl

from cheese_barcode import make_barcode, parse_weight_rows
from cheese_label_image import render_label_image

# --- EDIT THIS to the real path of the cheese label .btw template ---
CHEESE_LABEL_TEMPLATE_PATH = r"C:\Labels\CheeseLabel.btw"

# Optional header/divider label, same idea as HEADER_TEMPLATES in
# webhook_receiver.py. Leave as None to skip printing a header label.
CHEESE_HEADER_TEMPLATE_PATH = None  # e.g. r"C:\Labels\HeaderCheeseLabel.btw"

CHEESE_ITEMS_XLSX = r"C:\BarTenderWebhook\cheese_items.xlsx"

cheese_bp = Blueprint("cheese_labels", __name__)


# ---------------------------------------------------------------
# Items file (separate from ITEMS_EXCEL_PATH / test_items.xlsx used by
# the existing print.html tool — different schema, different purpose)
# ---------------------------------------------------------------

def load_cheese_items():
    wb = openpyxl.load_workbook(CHEESE_ITEMS_XLSX, read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = dict(zip(headers, row))
        if item.get("sku"):
            items.append(item)
    return items


def save_cheese_items(items):
    headers = ["sku", "productEn", "productHe", "itemCode", "kosher", "passover"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "cheese_items"
    ws.append(headers)
    for item in items:
        ws.append([item.get(h, "") for h in headers])
    wb.save(CHEESE_ITEMS_XLSX)


@cheese_bp.route("/cheese-items", methods=["GET"])
def cheese_items():
    """Feeds the product dropdown on cheese-labels.html."""
    items = load_cheese_items()
    return jsonify([
        {"sku": i["sku"], "productEn": i["productEn"], "productHe": i["productHe"]}
        for i in items
    ])


@cheese_bp.route("/cheese-items", methods=["POST"])
def add_cheese_item():
    """
    Adds a new product row to cheese_items.xlsx.
    Body (JSON): { sku, productEn, productHe, itemCode, kosher, passover }
    """
    data = request.get_json(force=True)
    sku = (data.get("sku") or "").strip()
    product_en = (data.get("productEn") or "").strip()
    product_he = (data.get("productHe") or "").strip()
    item_code = (data.get("itemCode") or "").strip()
    kosher = (data.get("kosher") or "").strip()
    passover = (data.get("passover") or "").strip()

    if not sku or not product_en or not product_he:
        return jsonify({"error": "יש למלא sku, שם באנגלית ושם בעברית"}), 400
    if not (item_code.isdigit() and len(item_code) == 4):
        return jsonify({"error": "קוד פריט חייב להיות בדיוק 4 ספרות"}), 400

    items = load_cheese_items()
    if any(str(i["sku"]) == sku for i in items):
        return jsonify({"error": f"SKU כבר קיים: {sku}"}), 400

    items.append({
        "sku": sku, "productEn": product_en, "productHe": product_he,
        "itemCode": item_code, "kosher": kosher, "passover": passover,
    })
    save_cheese_items(items)
    return jsonify({"ok": True, "sku": sku})


# ---------------------------------------------------------------
# Preview (no printing) — pasted text or uploaded Excel
# ---------------------------------------------------------------

def _rows_with_barcodes(parsed_rows, item_code):
    rows = []
    seen = {}
    for idx, pr in enumerate(parsed_rows, start=1):
        bc = make_barcode(item_code, pr["weightKg"])
        seen.setdefault(bc, []).append(idx)
        rows.append({"index": idx, "weightKg": pr["weightKg"], "uniqueId": pr["uniqueId"], "barcode": bc})
    for r in rows:
        r["duplicate"] = len(seen[r["barcode"]]) > 1
    return rows


@cheese_bp.route("/cheese-preview", methods=["POST"])
def cheese_preview():
    """
    Body (JSON): { "sku": "...", "weightsText": "4.34\\n4.39 1220\\n..." }
    Each line is a weight, optionally followed by a unique carton ID
    (e.g. "4.39 1220"). Returns the computed rows WITHOUT printing.
    Duplicate barcodes (same weight) are flagged for information only —
    NOT an error, and do not block printing.
    """
    data = request.get_json(force=True)
    sku = data.get("sku")
    parsed_rows = parse_weight_rows(data.get("weightsText", ""))

    items = {i["sku"]: i for i in load_cheese_items()}
    item = items.get(sku)
    if not item:
        return jsonify({"error": f"SKU not found: {sku}"}), 400
    if not parsed_rows:
        return jsonify({"error": "No weights provided"}), 400

    item_code = str(item["itemCode"]).zfill(4)
    rows = _rows_with_barcodes(parsed_rows, item_code)

    return jsonify({
        "sku": sku, "productHe": item["productHe"], "productEn": item["productEn"],
        "count": len(rows), "rows": rows,
    })


@cheese_bp.route("/cheese-preview-upload", methods=["POST"])
def cheese_preview_upload():
    """Same as /cheese-preview, but weights come from an uploaded Excel file
    (multipart/form-data: 'file' + 'sku'). Column A = weight, column B
    (optional) = unique carton ID."""
    sku = request.form.get("sku")
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "No file uploaded"}), 400

    wb = openpyxl.load_workbook(io.BytesIO(file.read()), read_only=True)
    ws = wb.active
    parsed_rows = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        cells = [c for c in row if c is not None]
        if not cells or not isinstance(cells[0], (int, float)):
            continue
        weight = float(cells[0])
        unique_id = str(cells[1]) if len(cells) > 1 else ""
        parsed_rows.append({"weightKg": weight, "uniqueId": unique_id})

    items = {i["sku"]: i for i in load_cheese_items()}
    item = items.get(sku)
    if not item:
        return jsonify({"error": f"SKU not found: {sku}"}), 400
    if not parsed_rows:
        return jsonify({"error": "No numeric weights found in file"}), 400

    item_code = str(item["itemCode"]).zfill(4)
    rows = _rows_with_barcodes(parsed_rows, item_code)

    return jsonify({
        "sku": sku, "productHe": item["productHe"], "productEn": item["productEn"],
        "count": len(rows), "rows": rows,
    })


@cheese_bp.route("/cheese-label-image", methods=["GET"])
def cheese_label_image():
    """Visual preview of ONE label (approximation, not the real .btw render).
    GET /cheese-label-image?sku=GOUDA48&weight=4.34&cheeseNo=1219"""
    sku = request.args.get("sku")
    weight = request.args.get("weight", type=float)
    cheese_no = request.args.get("cheeseNo", "")

    items = {i["sku"]: i for i in load_cheese_items()}
    item = items.get(sku)
    if not item or weight is None:
        return jsonify({"error": "sku and weight are required"}), 400

    item_code = str(item["itemCode"]).zfill(4)
    bc = make_barcode(item_code, weight)
    png_bytes = render_label_image(bc, weight, item["productHe"], item["productEn"], cheese_no)
    return Response(png_bytes, mimetype="image/png")


# ---------------------------------------------------------------
# Real BarTender printing — one open format, one PrintOut() per carton
# ---------------------------------------------------------------

def print_cheese_batch_via_bartender(rows, batch_number, product_en, product_he):
    """
    rows: list of {"weightKg": float, "uniqueId": str, "barcode": str}
    Opens the cheese label format ONCE, loops setting field values and
    calling PrintOut() per carton (NOT IdenticalCopiesOfLabel — every
    carton has different weight/barcode/id).
    """
    pythoncom.CoInitialize()
    bt = None
    try:
        bt = win32com.client.Dispatch("BarTender.Application")

        if CHEESE_HEADER_TEMPLATE_PATH:
            header_fmt = bt.Formats.Open(CHEESE_HEADER_TEMPLATE_PATH, False, "")
            try:
                header_fields = {
                    "Product": product_he,
                    "Quantity": str(len(rows)),
                    "BatchNumber": batch_number,
                }
                for fname, fvalue in header_fields.items():
                    try:
                        header_fmt.SetNamedSubStringValue(fname, fvalue)
                    except Exception as e:
                        logging.warning(f"שדה כותרת '{fname}' לא נמצא: {e}")
                header_fmt.PrintOut(False, False)
            finally:
                header_fmt.Close(2)

        fmt = bt.Formats.Open(CHEESE_LABEL_TEMPLATE_PATH, False, "")
        try:
            for r in rows:
                fields = {
                    "מספר_אצווה": batch_number,
                    "מזהה_קרטון": r.get("uniqueId", "") or "",
                    "ברקוד": r["barcode"],
                    "משקל": f'{r["weightKg"]:.2f}',
                    "שם_אנגלית": product_en,
                    "שם_עברית": product_he,
                }
                for field, value in fields.items():
                    try:
                        fmt.SetNamedSubStringValue(field, value)
                    except Exception as e:
                        logging.warning(f"שדה '{field}' לא נמצא ב-{CHEESE_LABEL_TEMPLATE_PATH}: {e}")
                fmt.PrintOut(False, False)
        finally:
            fmt.Close(2)

    finally:
        if bt is not None:
            try:
                bt.Quit(2)
            except Exception:
                pass
        pythoncom.CoUninitialize()


@cheese_bp.route("/print-cheese-batch", methods=["POST"])
def print_cheese_batch():
    """
    Body (JSON): { "sku": "...", "batchNumber": "...",
                   "rows": [{"weightKg":4.34,"uniqueId":"1219"}, ...] }
    Called after the user reviewed the preview and confirmed.
    """
    data = request.get_json(force=True)
    sku = data.get("sku")
    batch_number = data.get("batchNumber", "")
    rows_in = data.get("rows", [])

    items = {i["sku"]: i for i in load_cheese_items()}
    item = items.get(sku)
    if not item:
        return jsonify({"error": f"SKU not found: {sku}"}), 400
    if not rows_in:
        return jsonify({"error": "No rows provided"}), 400

    item_code = str(item["itemCode"]).zfill(4)
    rows = []
    for r in rows_in:
        w = r["weightKg"]
        rows.append({
            "weightKg": w,
            "uniqueId": r.get("uniqueId", ""),
            "barcode": make_barcode(item_code, w),
        })

    try:
        print_cheese_batch_via_bartender(rows, batch_number, item["productEn"], item["productHe"])
    except Exception as e:
        logging.error(f"הדפסת סבב גבינות נכשלה (sku={sku}): {e}")
        return jsonify({"error": str(e)}), 500

    logging.info(f"הודפסו {len(rows)} תוויות גבינה עבור {sku} (batch {batch_number})")
    return jsonify({"printed": len(rows), "sku": sku})


# In webhook_receiver.py, add near the bottom (before `if __name__ == "__main__":`):
#   from cheese_labels_routes import cheese_bp
#   app.register_blueprint(cheese_bp)
